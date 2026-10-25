"""The phase 2 fields on GET /deals/{id}/case (edits, adjustments, brief_draft) and the store that holds them.

The gateway's edit routes (Cayden's C2) do not exist yet, so these tests put a real Edit and a LearnedAdjustment
straight into the store. Once C2 lands, the same checks should pass through the gateway path.
Figures: Nu Aire "Net Income" is 5,438,271 USD in the fixture, 5,438.271 in USD thousands; the edit sets it to 6,438.271.
"""
import os

os.environ["AGENT_BACKEND"] = "scripted"
os.environ["AGENT_FALLBACK"] = ""

from io import BytesIO  # noqa: E402

import openpyxl  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.core.models import AgreementText, Edit, LearnedAdjustment  # noqa: E402
from app.db.fixtures.kewaunee import DEAL_ID, TARGET  # noqa: E402
from app.db.memory import MemoryStore  # noqa: E402
from app.main import create_app  # noqa: E402

REASON = "One-time acquisition cost, an allowed add-back"


@pytest.fixture
def app():
    return create_app()


@pytest.fixture
def client(app):
    return TestClient(app)


def case(c):
    r = c.get(f"/deals/{DEAL_ID}/case")
    assert r.status_code == 200
    return r.json()


def run_to_posted(c):
    assert c.post(f"/deals/{DEAL_ID}/run").json()["stage"] == "awaiting_flag"
    for f in case(c)["flags"]:
        assert c.post(f"/work-items/{f['id']}/answer?auto_advance=false", json={"option_id": f["default"]}).status_code == 200
    assert c.post(f"/deals/{DEAL_ID}/run").json()["stage"] == "awaiting_approval"
    item = case(c)["verdict"]["approval_item"]
    assert c.post(f"/work-items/{item}/approve").json()["state"] == "POSTED"


def put_edit_and_adjustment(store):
    """Stand-in for the gateway's edit route (C2, not built yet): write the records the route will write."""
    e = Edit(id=store.next_id("e"), deal_id=DEAL_ID, kind="amount", company=TARGET, line="Net Income",
             section="income_statement", old_value=5438.271, new_value=6438.271, status="reasoned",
             reason_text=REASON)
    a = LearnedAdjustment(id=store.next_id("a"), deal_id=DEAL_ID, company=TARGET, line="Net Income",
                          amount=6438.271, reason=REASON, edit_id=e.id)
    e = e.model_copy(update={"adjustment_id": a.id})
    store.edits[e.id] = e
    store.adjustments[a.id] = a
    return e, a


def test_case_has_empty_phase2_fields_before_and_after_a_run(client):
    c = case(client)
    assert c["edits"] == [] and c["adjustments"] == [] and c["brief_draft"] is None
    assert c["brief_verification"] is None
    run_to_posted(client)
    c = case(client)
    assert c["edits"] == [] and c["adjustments"] == []
    # The demo deal reads its Brief from the agreement, and the verifier approves it (D23).
    assert c["brief_draft"]["citations"] and c["deal"]["brief_status"] == "approved"
    v = c["brief_verification"]
    assert (v["passed"], v["failed"], v["auto_approved"]) == (22, 0, True)
    assert set(v["checks"][0]) == {"field", "value", "section", "quote", "ok", "reason", "checks"}


def test_case_export_and_certificate_show_a_stored_edit(app, client):
    run_to_posted(client)
    store = app.state.gateway.store
    e, a = put_edit_and_adjustment(store)
    assert (e.id, a.id) == ("e_001", "a_001")

    c = case(client)
    assert [x["id"] for x in c["edits"]] == ["e_001"]
    assert c["edits"][0]["new_value"] == 6438.271 and c["edits"][0]["adjustment_id"] == "a_001"
    assert c["adjustments"] == [a.model_dump(mode="json")]

    r = client.get(f"/deals/{DEAL_ID}/export.xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.content))
    ws = wb["Combined"]
    # Nu Aire "Net Income" maps onto Kewaunee "Net earnings": the target column (4) of that row is marked.
    row = next(i for i, v in enumerate(ws.iter_rows(values_only=True), start=1) if v[0] == "Net earnings")
    cell = ws.cell(row=row, column=4)
    assert cell.fill.fill_type == "solid"
    assert "Line was 5,438." in cell.comment.text and REASON in cell.comment.text
    edits_rows = [list(v) for v in wb["Edits"].iter_rows(values_only=True)]
    assert len(edits_rows) == 2 and edits_rows[1][:5] == ["Net Income", TARGET, 5438.271, 6438.271, REASON]

    md = client.get(f"/deals/{DEAL_ID}/certificate").text
    section = md.split("## Edits in review")[1].split("## ")[0]
    assert "e_001" in section and "Net Income" in section and REASON in section


def test_reverted_edit_stays_on_case_but_not_in_the_files(app, client):
    run_to_posted(client)
    store = app.state.gateway.store
    e, _ = put_edit_and_adjustment(store)
    store.edits[e.id] = e.model_copy(update={"status": "reverted"})
    assert [x["status"] for x in case(client)["edits"]] == ["reverted"]
    assert store.active_edits(DEAL_ID) == []
    cert = client.get(f"/deals/{DEAL_ID}/certificate?format=json").json()
    assert cert["edits"] == []


def test_reset_clears_edits_and_keeps_adjustments(app, client):
    run_to_posted(client)
    store = app.state.gateway.store
    _, a = put_edit_and_adjustment(store)
    assert client.post("/dev/reset").json()["ok"]
    c = case(client)
    assert c["edits"] == []
    assert [x["id"] for x in c["adjustments"]] == [a.id]
    # The adjustment id sequence survives with the memory, so the next one does not reuse a_001.
    assert store.next_id("a") == "a_002"
    assert client.post("/dev/reset?keep_memory=false").json()["ok"]
    assert case(client)["adjustments"] == []


def test_store_helpers_keep_order_and_drop_reverted():
    store = MemoryStore()
    base = dict(deal_id=DEAL_ID, kind="amount", company=TARGET, line="Net Income", section="income_statement",
                new_value=1.0)
    for eid, status in (("e_003", "needs_reason"), ("e_001", "reverted"), ("e_002", "needs_reason")):
        store.edits[eid] = Edit(id=eid, status=status, **base)
    store.edits["x"] = Edit(id="x", **{**base, "deal_id": "other"})
    assert [e.id for e in store.deal_edits(DEAL_ID)] == ["e_003", "e_001", "e_002"]
    assert [e.id for e in store.active_edits(DEAL_ID)] == ["e_003", "e_002"]
    store.agreements["ag_1"] = AgreementText(id="ag_1", deal_id=DEAL_ID, text="t", sha256="0" * 64, source="8-K")
    store.reset_deals(keep_memory=True)
    assert store.edits == {} and store.agreements == {}


def test_openapi_case_schema_has_phase2_fields(client):
    schemas = client.get("/openapi.json").json()["components"]["schemas"]
    props = schemas["CaseResponse"]["properties"]
    assert {"edits", "adjustments", "brief_draft"} <= set(props)
    assert {"Edit", "LearnedAdjustment", "BriefDraft"} <= set(schemas)
