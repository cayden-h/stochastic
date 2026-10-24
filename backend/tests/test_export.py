"""Final files: GET /deals/{id}/export.xlsx and the certificate's "Edits in review". Figures come from the tool runs."""
import os

os.environ["AGENT_BACKEND"] = "scripted"
os.environ["AGENT_FALLBACK"] = ""

from datetime import datetime, timezone  # noqa: E402
from io import BytesIO  # noqa: E402
from types import SimpleNamespace  # noqa: E402

import openpyxl  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.api import certificate as cert_doc  # noqa: E402
from app.api.export import AMBER, build_workbook, export_filename, workbook_bytes  # noqa: E402
from app.db.fixtures.kewaunee import DEAL_ID, PLATFORM, TARGET  # noqa: E402
from app.main import create_app  # noqa: E402

URL = f"/deals/{DEAL_ID}/export.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GATE = "final files are available after the controller approves the result"


def case(c):
    return c.get(f"/deals/{DEAL_ID}/case").json()


def run_to_proposed(c):
    assert c.post(f"/deals/{DEAL_ID}/run").json()["stage"] == "awaiting_flag"
    for f in case(c)["flags"]:
        assert c.post(f"/work-items/{f['id']}/answer?auto_advance=false", json={"option_id": f["default"]}).status_code == 200
    assert c.post(f"/deals/{DEAL_ID}/run").json()["stage"] == "awaiting_approval"


def run_to_posted(c):
    run_to_proposed(c)
    item = case(c)["verdict"]["approval_item"]
    assert c.post(f"/work-items/{item}/approve").json()["state"] == "POSTED"


@pytest.fixture
def app():
    return create_app()


@pytest.fixture
def client(app):
    return TestClient(app)


def sheet_rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def load(content: bytes):
    return openpyxl.load_workbook(BytesIO(content))


def edit(**kw):
    base = dict(id="ed_1", deal_id=DEAL_ID, kind="amount", company=TARGET, line="Net Income",
                section="income_statement", old_value=5438.271, new_value=6438.271, old_platform_line=None,
                new_platform_line=None, status="reasoned", reason_chip=None,
                reason_text="One-time acquisition cost, an allowed add-back", remember=True, policy_id="p_12",
                adjustment_id=None, actor="human:controller",
                at=datetime(2026, 10, 25, 10, 42, tzinfo=timezone.utc), run_ids=[])
    base.update(kw)
    return SimpleNamespace(**base)


def posted_parts(app, client):
    run_to_posted(client)
    gw = app.state.gateway
    store = gw.store
    deal = store.deals[DEAL_ID]
    result = store.latest_result(DEAL_ID)
    items = store.deal_items(DEAL_ID)
    return store, deal, result, items, gw.verify_chain()


# ---------------------------------------------------------------------------- the route


def test_404_for_unknown_deal(client):
    assert client.get("/deals/nope/export.xlsx").status_code == 404


def test_409_before_any_run(client):
    r = client.get(URL)
    assert r.status_code == 409 and r.json()["detail"] == GATE


def test_409_after_proposed_before_approval(client):
    run_to_proposed(client)
    r = client.get(URL)
    assert r.status_code == 409 and r.json()["detail"] == GATE


def test_200_after_approval_with_type_and_filename(client):
    run_to_posted(client)
    r = client.get(URL)
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"] == f'attachment; filename="{PLATFORM} + {TARGET} combined books.xlsx"'
    wb = load(r.content)
    assert wb.sheetnames == ["Combined", "Edits", "Audit"]


def test_filename_is_ascii_only():
    name = export_filename("Kéwaunee \"Sci\"", "Nu\u2013Aire")
    assert name == "K waunee  Sci  + Nu Aire combined books.xlsx"
    assert name.isascii()


def test_combined_sheet_equals_result_rows(client):
    run_to_posted(client)
    result = case(client)["result"]
    rows = sheet_rows(load(client.get(URL).content)["Combined"])
    assert rows[0] == ["Line", "Section", PLATFORM, TARGET, "Combined", "Run id"]
    body = rows[1:]
    assert len(body) == len(result["rows"])
    for got, want in zip(body, result["rows"]):
        assert got == [want["line"], want["section"], want["platform"], want["target"], want["combined"],
                       want["run_id"]]
        assert all(isinstance(x, (int, float)) for x in got[2:5])  # numbers, not strings
    total = next(r for r in body if r[0] == "Total Assets")
    assert total[4] == pytest.approx(182456.181, abs=0.001)  # hand-computed, test_demo_run.py


LIMIT_KEYS = {"Limit test leverage": "leverage", "Limit test limit": "limit", "Limit test verdict": "verdict",
              "Limit test headroom (turns)": "headroom_turns", "Limit test debt capacity": "debt_capacity"}
FIGURE_KEYS = {"EBITDA": ("ebitda", "ebitda"), "Debt": ("covenant_debt", "debt"),
               "Leverage": ("leverage", "leverage"), "Tie-out ok": ("tie_out", "ok"),
               **{k: ("limit_test", v) for k, v in LIMIT_KEYS.items()}}


def test_every_audit_figure_equals_its_run(client):
    run_to_posted(client)
    rows = sheet_rows(load(client.get(URL).content)["Audit"])
    assert rows[0] == ["Figure", "Value", "Run id", "Tool"]
    figures = {r[0]: r for r in rows[1:] if r[0] in FIGURE_KEYS}
    assert set(figures) == set(FIGURE_KEYS)
    for name, (tool, key) in FIGURE_KEYS.items():
        _, value, run_id, got_tool = figures[name]
        run = client.get(f"/runs/{run_id}").json()
        assert run["tool"] == tool == got_tool, name
        assert value == run["outputs"][key], name
    assert figures["EBITDA"][1] == pytest.approx(28378.845, abs=0.001)
    assert figures["Debt"][1] == 21633
    assert figures["Leverage"][1] == pytest.approx(0.762, abs=0.0005)
    assert figures["Limit test debt capacity"][1] == pytest.approx(49314.113, abs=0.001)
    assert figures["Limit test verdict"][1] == "ok" and figures["Tie-out ok"][1] is True

    result = case(client)["result"]
    combined = [r for r in rows[1:] if r[0] and str(r[0]).startswith("Combined: ")]
    assert len(combined) == len(result["rows"])
    for got, want in zip(combined, result["rows"]):
        assert got[0] == f"Combined: {want['line']} ({want['section']})" and got[2] == want["run_id"]
        run = client.get(f"/runs/{got[2]}").json()
        assert run["tool"] == got[3] == "consolidate"
        out = next(x for x in run["outputs"]["rows"] if x["line"] == want["line"] and x["section"] == want["section"])
        assert got[1] == out["combined"]

    tail = [r for r in rows if r[0] is not None][-6:]
    audit = client.get(f"/deals/{DEAL_ID}/audit").json()
    assert ["Audit rows", len(audit["rows"]), None, None] in tail
    assert ["Chain valid", True, None, None] in tail
    assert ["Denied attempts", 0, None, None] in tail
    assert any(r[0] == "Backend" and r[1] == "scripted" for r in tail)
    assert tail[-1][0] == "Draft. Not signed."
    assert [None, None, None, None] in rows  # the blank row before the summary


def test_edits_sheet_is_header_only_without_edits(client):
    run_to_posted(client)
    rows = sheet_rows(load(client.get(URL).content)["Edits"])
    assert rows == [["Line", "Company", "Agent value", "Final value", "Reason", "Who", "When", "Policy id"]]


def test_certificate_says_no_edits(client):
    run_to_posted(client)
    md = client.get(f"/deals/{DEAL_ID}/certificate").text
    assert "## Edits in review\n\n- None." in md
    assert client.get(f"/deals/{DEAL_ID}/certificate?format=json").json()["edits"] == []


# ---------------------------------------------------------------------------- build_workbook with edits


def test_reasoned_edit_marks_the_cell_and_lists_the_edit(app, client):
    store, deal, result, items, chain = posted_parts(app, client)
    wb = load(workbook_bytes(build_workbook(store, deal, result, items, chain, edits=[edit()])))
    ws = wb["Combined"]
    # Nu Aire "Net Income" maps to Kewaunee "Net earnings": the target cell on that row is the edited one.
    r = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1) if row[0] == "Net earnings")
    cell = ws.cell(row=r, column=4)
    assert cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb.endswith(AMBER)
    text = cell.comment.text
    assert text == ("Edited by human:controller, 2026-10-25 10:42. Was 5,438.271. "
                    "Reason: One-time acquisition cost, an allowed add-back. Policy p_12.")
    assert ws.cell(row=r, column=3).comment is None  # the platform cell is not edited
    unedited = ws.cell(row=r + 1, column=4)
    assert unedited.comment is None and unedited.fill.fill_type is None

    rows = sheet_rows(wb["Edits"])
    assert len(rows) == 2
    assert rows[1] == ["Net Income", TARGET, 5438.271, 6438.271, "One-time acquisition cost, an allowed add-back",
                       "human:controller", "2026-10-25 10:42", "p_12"]


def test_platform_edit_with_chip_and_no_policy(app, client):
    store, deal, result, items, chain = posted_parts(app, client)
    e = edit(company=PLATFORM, line="Net earnings", old_value=19057.0, new_value=19500.0, reason_text=None,
             reason_chip=1, policy_id=None)
    ws = load(workbook_bytes(build_workbook(store, deal, result, items, chain, edits=[e])))["Combined"]
    r = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1) if row[0] == "Net earnings")
    cell = ws.cell(row=r, column=3)
    assert cell.fill.fill_type == "solid"
    assert cell.comment.text == "Edited by human:controller, 2026-10-25 10:42. Was 19,057.000. Reason: chip 1."
    assert ws.cell(row=r, column=4).comment is None


def test_reverted_edit_marks_nothing(app, client):
    store, deal, result, items, chain = posted_parts(app, client)
    wb = load(workbook_bytes(build_workbook(store, deal, result, items, chain, edits=[edit(status="reverted")])))
    ws = wb["Combined"]
    for row in ws.iter_rows():
        for cell in row:
            assert cell.comment is None and cell.fill.fill_type is None
    assert len(sheet_rows(wb["Edits"])) == 1


def test_edits_default_reads_the_store(app, client):
    store, deal, result, items, chain = posted_parts(app, client)
    store.edits = {"ed_1": edit(), "ed_2": edit(id="ed_2", deal_id="other_deal"), "ed_3": edit(id="ed_3", status="reverted")}
    try:
        wb = load(workbook_bytes(build_workbook(store, deal, result, items, chain)))
        assert len(sheet_rows(wb["Edits"])) == 2
        cert = cert_doc.build_certificate(store, deal, result, items, chain)
        assert [x["id"] for x in cert["edits"]] == ["ed_1"]
    finally:
        del store.edits


def test_certificate_lists_each_edit_and_its_reason(app, client):
    store, deal, result, items, chain = posted_parts(app, client)
    remap = edit(id="ed_2", kind="remap", line="Certificates of Deposit", section="balance_sheet", old_value=None,
                 new_value=None, old_platform_line="Cash and cash equivalents",
                 new_platform_line="Prepaid expenses and other current assets", reason_text="Not a cash equivalent",
                 policy_id=None)
    cert = cert_doc.build_certificate(store, deal, result, items, chain, edits=[edit(), remap, edit(id="ed_3", status="reverted")])
    assert cert["edits"][0] == {"id": "ed_1", "kind": "amount", "company": TARGET, "line": "Net Income",
                                "old": 5438.271, "new": 6438.271,
                                "reason": "One-time acquisition cost, an allowed add-back",
                                "actor": "human:controller", "at": "2026-10-25T10:42:00+00:00", "policy_id": "p_12"}
    assert cert["edits"][1]["old"] == "Cash and cash equivalents"
    assert len(cert["edits"]) == 2
    md = cert_doc.render_markdown(cert)
    section = md.split("## Edits in review")[1].split("## ")[0]
    assert "ed_1" in section and "One-time acquisition cost, an allowed add-back" in section and "Policy p_12" in section
    assert "5,438.271 to 6,438.271" in section
    assert "Cash and cash equivalents to Prepaid expenses and other current assets" in section
    assert "Not a cash equivalent" in section and "ed_3" not in section
    assert "Edits in review" in cert_doc.render_html(cert)

    # the remap marks the target cell on the new platform line
    ws = load(workbook_bytes(build_workbook(store, deal, result, items, chain, edits=[remap])))["Combined"]
    r = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1)
             if row[0] == "Prepaid expenses and other current assets")
    assert "Was Cash and cash equivalents." in ws.cell(row=r, column=4).comment.text
