"""The edit routes (docs/tasks/tri.md T2) through HTTP: the demo beat of design/review-companion.md section 9.

Thin routes over the gateway (apply_edit, reason_edit, revert_edit). Every figure is read from the recompute's tool
runs. Expected values are worked out by hand from the fixture (app/db/fixtures/kewaunee.py):
Nu Aire General and Administrative 6,162,806 dollars = 6,162.806; cut 500 to 5,662.806. By the income statement
roll-up rule Net Income moves +500, so covenant EBITDA 28,378.845 + 500 = 28,878.845 and leverage
21,633 / 28,878.845 = 0.74909 -> 0.749 (the agents' figure is 21,633 / 28,378.845 = 0.762).
"""
import json
import os
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient

os.environ["AGENT_BACKEND"] = "scripted"
os.environ["AGENT_FALLBACK"] = "scripted"

from app.api.events import next_events  # noqa: E402
from app.api.export import AMBER  # noqa: E402
from app.api.schemas import CaseResponse  # noqa: E402
from app.core.models import MAPPER, PREPARER, REVIEWER  # noqa: E402
from app.db.fixtures.kewaunee import DEAL_ID, PLATFORM, TARGET  # noqa: E402
from app.main import create_app  # noqa: E402

LINE = "General and Administrative"
AMOUNT = {"kind": "amount", "company": TARGET, "line": LINE, "value": 5662.806, "section": "income_statement"}
CHIP_0 = "One-time acquisition cost, an allowed add-back"


def case(c):
    return c.get(f"/deals/{DEAL_ID}/case").json()


@pytest.fixture
def c():
    """A fresh app, run to the controller's approval with every flag answered by its default."""
    client = TestClient(create_app())
    assert client.post(f"/deals/{DEAL_ID}/run").json()["stage"] == "awaiting_flag"
    for f in case(client)["flags"]:
        r = client.post(f"/work-items/{f['id']}/answer?auto_advance=false", json={"option_id": f["default"]})
        assert r.status_code == 200
    assert client.post(f"/deals/{DEAL_ID}/run").json()["stage"] == "awaiting_approval"
    assert case(client)["verdict"]["leverage"] == 0.762
    return client


def parse(events):
    out = []
    for e in events:
        lines = e.strip().split("\n")
        out.append((lines[0].removeprefix("event: "), json.loads(lines[1].removeprefix("data: "))))
    return out


def test_the_demo_beat_through_http(c):
    r = c.post(f"/deals/{DEAL_ID}/edits", json=AMOUNT)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["verdict"]["ebitda"] == 28878.845 and out["verdict"]["leverage"] == 0.749
    assert out["verdict"]["verdict"] == "ok" and out["verdict"]["debt"] == 21633.0
    edit = out["edit"]
    assert edit["status"] == "needs_reason" and (edit["old_value"], edit["new_value"]) == (6162.806, 5662.806)
    assert out["reasons"][0] == {"id": 0, "label": CHIP_0}
    # every figure is the recompute's own run
    runs = out["result"]["run_ids_by_tool"]
    assert c.get(f"/runs/{runs['ebitda']}").json()["outputs"]["ebitda"] == 28878.845
    assert c.get(f"/runs/{runs['ebitda']}").json()["inputs"]["edit_ids"] == [edit["id"]]
    rollup = c.get(f"/runs/{runs['consolidate']}").json()["outputs"]["rollup"]
    assert [(x["line"], x["delta"]) for x in rollup] == [("Total Operating Expense", -500.0),
                                                         ("Income from Operations", 500.0),
                                                         ("Income Before Income Taxes", 500.0), ("Net Income", 500.0)]
    body = case(c)
    CaseResponse.model_validate(body)
    assert [e["id"] for e in body["edits"]] == [edit["id"]] and body["verdict"]["leverage"] == 0.749
    assert body["reason_chips"]["amount"][0]["label"] == CHIP_0 and len(body["reason_chips"]["remap"]) == 3

    approval = out["approval"]["id"]
    refused = c.post(f"/work-items/{approval}/approve")
    assert refused.status_code == 403 and edit["id"] in refused.json()["detail"]

    reasoned = c.post(f"/edits/{edit['id']}/reason", json={"chip": 0})
    assert reasoned.status_code == 200
    assert reasoned.json()["status"] == "reasoned" and reasoned.json()["adjustment_id"]
    assert reasoned.json()["reason_text"] == CHIP_0

    assert c.post(f"/work-items/{approval}/approve").json()["state"] == "POSTED"
    cert = c.get(f"/deals/{DEAL_ID}/certificate?format=json").json()
    assert [(e["id"], e["line"], e["old"], e["new"]) for e in cert["edits"]] == [
        (edit["id"], LINE, 6162.806, 5662.806)]
    assert cert["limit_test"]["leverage"] == 0.749 and cert["limit_test"]["verdict"] == "ok"
    assert cert["ebitda"]["ebitda"] == 28878.845

    xlsx = c.get(f"/deals/{DEAL_ID}/export.xlsx")
    assert xlsx.status_code == 200
    ws = openpyxl.load_workbook(BytesIO(xlsx.content))["Combined"]
    # the row General and Administrative lands on, found through the mappings as the export does
    lines = {l["id"]: l for l in body["source_lines"]}
    m = next(m for m in body["mappings"] if lines[m["target_line_id"]]["line"] == LINE)
    row = next(i for i, v in enumerate(ws.iter_rows(values_only=True), start=1)
               if v[0] == m["platform_line"] and v[1] == "income_statement")
    cell = ws.cell(row=row, column=4)
    assert cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb.endswith(AMBER)
    assert "Line was 6,163" in cell.comment.text


def test_undo_restores_the_agents_leverage(c):
    agents = case(c)["result"]
    edit = c.post(f"/deals/{DEAL_ID}/edits", json=AMOUNT).json()["edit"]
    r = c.delete(f"/edits/{edit['id']}")
    assert r.status_code == 200
    out = r.json()
    assert out["edit"]["status"] == "reverted" and out["result"]["edit_ids"] == []
    assert (out["verdict"]["leverage"], out["verdict"]["ebitda"]) == (0.762, 28378.845)
    # every row of the revision equals the agents' row (the run id is the recompute's own), and so do the figures
    def rows(result):
        return [{k: v for k, v in row.items() if k != "run_id"} for row in result["rows"]]
    assert rows(out["result"]) == rows(agents) and len(rows(agents)) > 0
    assert {row["run_id"] for row in out["result"]["rows"]} != {row["run_id"] for row in agents["rows"]}
    for k in ("ebitda", "debt", "leverage", "limit", "verdict", "headroom_turns", "tie_out_ok"):
        assert out["result"][k] == agents[k], k
    assert case(c)["verdict"]["leverage"] == 0.762
    assert c.delete(f"/edits/{edit['id']}").status_code == 409  # already reverted
    assert c.delete("/edits/e_999").status_code == 404


def test_an_agent_actor_is_refused(c):
    for actor in (MAPPER, PREPARER, REVIEWER):
        r = c.post(f"/deals/{DEAL_ID}/edits", json=AMOUNT, headers={"X-Actor": actor})
        assert r.status_code == 403 and "only human:controller" in r.json()["detail"]
    assert case(c)["edits"] == []
    edit = c.post(f"/deals/{DEAL_ID}/edits", json=AMOUNT).json()["edit"]
    assert c.post(f"/edits/{edit['id']}/reason", json={"chip": 0}, headers={"X-Actor": PREPARER}).status_code == 403
    assert c.delete(f"/edits/{edit['id']}", headers={"X-Actor": REVIEWER}).status_code == 403


def detail(r):
    return json.dumps(r.json()["detail"])


def test_bad_bodies_get_422_naming_the_field(c):
    url = f"/deals/{DEAL_ID}/edits"
    no_value = {k: v for k, v in AMOUNT.items() if k != "value"}
    r = c.post(url, json=no_value)
    assert r.status_code == 422 and "value is required" in detail(r)
    for bad in ("NaN", "1e999", "-Infinity"):  # Python's JSON reader accepts these; the route must not
        r = c.post(url, content='{"kind": "amount", "company": "Nu Aire", "line": "General and Administrative", '
                                f'"value": {bad}}}', headers={"Content-Type": "application/json"})
        assert r.status_code == 422 and "value must be a finite number" in detail(r), bad
        assert r.json()["detail"][0]["loc"] == ["body", "value"]
        assert r.json()["detail"][0]["input"] == "non-finite"  # never echoed as a non-finite number
    # a boolean is not a number, and a string is not a number either, even "NaN"
    for bad in (True, "NaN", "12"):
        r = c.post(url, json={**AMOUNT, "value": bad})
        assert r.status_code == 422 and r.json()["detail"][0]["loc"] == ["body", "value"], bad
    r = c.post(url, json={"kind": "remap", "company": TARGET, "line": "Certificates of Deposit"})
    assert r.status_code == 422 and "platform_line is required" in detail(r)
    r = c.post(url, json={**AMOUNT, "kind": "split"})
    assert r.status_code == 422 and "kind" in detail(r)
    assert c.post("/deals/nope/edits", json=AMOUNT).status_code == 404
    # the gateway's own refusals carry their status through: a subtotal, a line the company does not have
    r = c.post(url, json={**AMOUNT, "line": "Net Income"})
    assert r.status_code == 422 and "subtotal" in r.json()["detail"]
    r = c.post(url, json={**AMOUNT, "company": PLATFORM, "line": "No Such Line"})
    assert r.status_code == 422
    edit = c.post(url, json=AMOUNT).json()["edit"]
    r = c.post(f"/edits/{edit['id']}/reason", json={"remember": True})
    assert r.status_code == 422 and "chip or text is required" in detail(r)
    r = c.post(f"/edits/{edit['id']}/reason", json={"text": "   "})
    assert r.status_code == 422 and "chip or text" in detail(r)
    assert c.post(f"/edits/{edit['id']}/reason", json={"chip": 7}).status_code == 422
    r = c.post(f"/edits/{edit['id']}/reason", json={"chip": True})
    assert r.status_code == 422 and r.json()["detail"][0]["loc"] == ["body", "chip"]
    assert c.post("/edits/e_999/reason", json={"chip": 0}).status_code == 404
    typed = c.post(f"/edits/{edit['id']}/reason", json={"text": "Bonus accrual booked twice", "remember": False})
    assert typed.status_code == 200 and typed.json()["adjustment_id"] is None


def test_non_finite_numbers_in_any_field_get_422_naming_the_field(c):
    """JSON NaN anywhere in a body is a 422 that names the field (app/api/errors.py), never a 500."""
    edit = c.post(f"/deals/{DEAL_ID}/edits", json=AMOUNT).json()["edit"]
    raw = '{"kind": "amount", "company": "Nu Aire", "line": "General and Administrative", "value": 1.0, '
    for url, body, loc in (
            (f"/edits/{edit['id']}/reason", '{"chip": NaN}', ["body", "chip"]),
            (f"/deals/{DEAL_ID}/edits", raw + '"section": NaN}', ["body", "section"]),
            (f"/deals/{DEAL_ID}/edits", raw.replace('"amount"', "NaN") + '"section": "income_statement"}',
             ["body", "kind"]),
            (f"/deals/{DEAL_ID}/edits", "[NaN]", ["body"])):
        r = c.post(url, content=body, headers={"Content-Type": "application/json"})
        assert r.status_code == 422, (url, body, r.text)
        first = r.json()["detail"][0]
        assert first["loc"] == loc and first["msg"], body
        assert "NaN" not in r.text and "Infinity" not in r.text
    assert len(case(c)["edits"]) == 1  # nothing else was kept


def test_a_remap_through_http_ties_out(c):
    r = c.post(f"/deals/{DEAL_ID}/edits", json={"kind": "remap", "company": TARGET, "line": "Certificates of Deposit",
                                                "platform_line": "Prepaid expenses and other current assets"})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["edit"]["new_platform_line"] == "Prepaid expenses and other current assets"
    assert out["verdict"]["tie_out_ok"] is True and out["verdict"]["leverage"] == 0.762  # a remap moves no total
    assert len(out["reasons"]) == 3


RD_REMAP = {"kind": "remap", "company": TARGET, "line": "Research and Development", "platform_line": "Cost of products sold"}


def figures(result, section="income_statement"):
    return {r["line"]: (r["target"], r["combined"]) for r in result["rows"] if r["section"] == section}


def test_a_remap_into_cost_of_sales_recomputes_gross_profit_through_http(c):
    """The console's reclass: R&D (532.092) moves from Operating expenses to Cost of products sold."""
    agents = case(c)["result"]
    r = c.post(f"/deals/{DEAL_ID}/edits", json=RD_REMAP)
    assert r.status_code == 200, r.text
    out = r.json()
    was, now = figures(agents), figures(out["result"])
    assert now["Gross profit"] == (23149.519, 75200.519)  # 23,681.611 - 532.092; 75,732.611 - 532.092
    assert now["Cost of products sold"] == (51700.783, 203404.783)
    assert now["Operating expenses"] == (17900.909, 51670.909)
    for line in ("Operating profit", "Profit before income taxes", "Net earnings"):
        assert now[line] == was[line], line  # a reclass within operating costs
    v = out["verdict"]
    assert (v["ebitda"], v["debt"], v["leverage"], v["tie_out_ok"]) == (28378.845, 21633.0, 0.762, True)
    # the figure is the recompute's consolidate run, which names the edit and the moved subtotal
    run = c.get(f"/runs/{out['result']['run_ids_by_tool']['consolidate']}").json()
    assert run["outputs"]["edits_applied"] == [out["edit"]["id"]]
    assert run["outputs"]["remap_rollup"] == [{"edit_id": out["edit"]["id"], "line": "Gross profit",
                                               "section": "income_statement", "delta": -532.092}]
    assert figures(case(c)["result"])["Gross profit"] == (23149.519, 75200.519)


def test_a_cost_cut_moves_every_displayed_subtotal_through_http(c):
    agents = figures(case(c)["result"])
    now = figures(c.post(f"/deals/{DEAL_ID}/edits", json=AMOUNT).json()["result"])
    for line in ("Operating profit", "Profit before income taxes", "Net earnings"):
        assert (round(now[line][0] - agents[line][0], 3), round(now[line][1] - agents[line][1], 3)) == (500, 500), line
    assert now["Gross profit"] == agents["Gross profit"]
    assert now["Net earnings"] == (5938.271, 24995.271)  # 5,438.271 + 500; 24,495.271 + 500


def test_a_balance_sheet_remap_that_breaks_the_tie_out_says_so(c):
    """Certificates of Deposit (an asset) onto Accounts payable (a liability): assets 991.082 lower, liabilities and
    equity 991.082 higher, so the console shows "Does not tie out" and the approval needs info."""
    r = c.post(f"/deals/{DEAL_ID}/edits", json={"kind": "remap", "company": TARGET, "line": "Certificates of Deposit",
                                                "platform_line": "Accounts payable"})
    assert r.status_code == 200, r.text
    out = r.json()
    bs = figures(out["result"], "balance_sheet")
    assert bs["Total Assets"] == (46987.099, 181465.099)
    assert bs["Total Liabilities and Stockholders' Equity"] == (48969.263, 183447.263)
    assert out["verdict"]["tie_out_ok"] is False and out["result"]["tie_out_ok"] is False
    assert out["approval"]["state"] == "NEEDS_INFO"


def test_undo_of_every_kind_of_edit_restores_the_agents_books(c):
    agents = case(c)["result"]
    ids = []
    for body in (RD_REMAP, AMOUNT, {"kind": "remap", "company": TARGET, "line": "Certificates of Deposit",
                                    "platform_line": "Accounts payable"}):
        r = c.post(f"/deals/{DEAL_ID}/edits", json=body)
        assert r.status_code == 200, r.text
        ids.append(r.json()["edit"]["id"])
    assert case(c)["verdict"]["tie_out_ok"] is False
    for eid in reversed(ids):
        out = c.delete(f"/edits/{eid}").json()
    assert out["result"]["edit_ids"] == []

    def rows(result):
        return [{k: v for k, v in row.items() if k != "run_id"} for row in result["rows"]]
    assert rows(out["result"]) == rows(agents)
    for k in ("ebitda", "debt", "leverage", "verdict", "tie_out_ok"):
        assert out["result"][k] == agents[k], k


def test_the_edit_stream_event():
    from app.agents.jobs import advance
    from app.core.models import CONTROLLER
    from app.db.fixtures.kewaunee import seed
    from app.db.memory import MemoryStore
    from app.gateway.gateway import Gateway

    store = MemoryStore()
    gw = Gateway(store)
    seed(store, gw)
    advance(gw, DEAL_ID)
    for f in gw.open_flags(DEAL_ID):
        gw.answer(CONTROLLER, f.id, f.default)
    advance(gw, DEAL_ID)
    cursor: dict = {}
    events, seen, last = next_events(store, DEAL_ID, 0, None, edits_seen=cursor)
    assert "edit" not in [k for k, _ in parse(events)]  # no edits yet: nothing to send
    # without the cursor the old behaviour holds: no edit event at all
    edit = gw.apply_edit(CONTROLLER, DEAL_ID, "amount", TARGET, LINE, value=5662.806)["edit"]
    plain, _, _ = next_events(store, DEAL_ID, seen, last)
    assert "edit" not in [k for k, _ in parse(plain)]
    events, seen, last = next_events(store, DEAL_ID, seen, last, edits_seen=cursor)
    sent = [d for k, d in parse(events) if k == "edit"]
    assert len(sent) == 1 and sent[0]["deal_id"] == DEAL_ID and sent[0]["at"]
    assert sent[0]["edits"] == [{"id": edit.id, "status": "needs_reason"}]
    events, seen, last = next_events(store, DEAL_ID, seen, last, edits_seen=cursor)
    assert "edit" not in [k for k, _ in parse(events)]  # unchanged: not sent again
    gw.reason_edit(CONTROLLER, edit.id, chip=0)
    events, seen, last = next_events(store, DEAL_ID, seen, last, edits_seen=cursor)
    assert [d["edits"] for k, d in parse(events) if k == "edit"] == [[{"id": edit.id, "status": "reasoned"}]]
    gw.revert_edit(CONTROLLER, edit.id)
    events, seen, last = next_events(store, DEAL_ID, seen, last, edits_seen=cursor)
    assert [d["edits"] for k, d in parse(events) if k == "edit"] == [[{"id": edit.id, "status": "reverted"}]]
