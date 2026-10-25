"""Final files: the combined books as an Excel workbook. Pure functions, no FastAPI.

Three sheets: Combined (the result rows), Edits (one row per active controller edit), Audit (the run id behind every
figure). Figures are read from the certificate, which reads them from the tool runs, so the two files agree.
Nothing here computes a deal number.
Owner: Tri (backend/app/api/).
"""
from __future__ import annotations

from datetime import datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from app.api import certificate as cert_doc
from app.tools.edits import _line_matches, apply_remap_edits

AMBER = "FBEBC8"
# Whole thousands with separators, negatives in parentheses: how the console shows a figure. A display format only;
# the cell keeps the tool's value.
THOUSANDS = "#,##0;(#,##0)"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EDITS_HEADER = ["Line", "Company", "Agent value", "Final value", "Reason", "Who", "When", "Policy id"]
AUDIT_HEADER = ["Figure", "Value", "Run id", "Tool"]


TARGET_EDIT_NOTE = ("A target edit is marked on the platform row its source line maps onto, after any remap; "
                    "the note gives the source line's own value, not the row total.")


def _norm(text) -> str:
    return (text or "").strip().lower()


def _when(at) -> str | None:
    """The time in UTC, labelled, so Excel does not read it as local time. A naive time is taken as UTC."""
    if isinstance(at, datetime):
        utc = at.astimezone(timezone.utc) if at.tzinfo else at
        return f"{utc:%Y-%m-%d %H:%M} UTC"
    return str(at) if at is not None else None


def export_filename(platform: str, target: str) -> str:
    """'<platform> + <target> combined books.xlsx', ASCII only; any other character (and quotes) becomes a space."""
    name = f"{platform} + {target} combined books.xlsx"
    return "".join(c if c.isascii() and c.isprintable() and c not in '"\\' else " " for c in name)


def whole(x: float) -> str:
    """A figure as the cells show it: whole thousands with separators, half up, a negative in parentheses."""
    v = int(Decimal(str(x)).quantize(Decimal(1), rounding=ROUND_HALF_UP))
    return f"({-v:,})" if v < 0 else f"{v:,}"


def edit_note(edit) -> str:
    """The Excel note on an edited cell. For an amount edit the old figure is the edited line's value."""
    old, _ = cert_doc.edit_values(edit)
    if getattr(edit, "kind", None) == "remap":
        was = f"Line was mapped to {old or 'n/a'}"
    else:
        num = whole(old) if isinstance(old, (int, float)) and not isinstance(old, bool) else (old or "n/a")
        was = f"Line was {num}"
    text = (f"Edited by {getattr(edit, 'actor', None) or 'unknown'}, {_when(getattr(edit, 'at', None))}. "
            f"{was}. Reason: {cert_doc.edit_reason(edit)}.")
    policy = getattr(edit, "policy_id", None)
    return text + (f" Policy {policy}." if policy else "")


def _platform_row_of(platform_lines, mapping, target_line):
    """(name, section) of the platform row a mapping lands on: by id, else by name in the target line's section
    first and then any section (the gateway's order)."""
    by_id = {p.id: p for p in platform_lines}
    p = by_id.get(mapping.platform_line_id) if mapping.platform_line_id else None
    if p is None and mapping.platform_line:
        key = _norm(mapping.platform_line)
        named = [x for x in platform_lines if _norm(x.line) == key]
        p = next((x for x in named if x.section == target_line.section), named[0] if named else None)
    return (_norm(p.line), p.section) if p else None


def _edited_cells(store, deal, rows, edits) -> dict[tuple[int, int], object]:
    """(row index into rows, column 3 platform or 4 target) -> the edit that marks that cell. The last edit wins.

    A platform edit marks the row of the same line name (and section, when the edit gives one). A target edit marks
    the row its source line maps onto after the active remaps are applied (apply_remap_edits), in the platform
    line's own section, since the Combined sheet has the platform's line names.
    """
    platform, target = _norm(deal.platform), _norm(deal.target)
    platform_lines = store.lines(deal.id, deal.platform)
    target_lines = store.lines(deal.id, deal.target)
    usable = [m for m in store.deal_mappings(deal.id) if not m.flagged or m.approved_by]
    mappings, _ = apply_remap_edits(usable, target_lines, platform_lines, edits)
    by_target = {m.target_line_id: m for m in mappings}
    out: dict[tuple[int, int], object] = {}
    for e in edits:
        company, section = _norm(getattr(e, "company", None)), getattr(e, "section", None)
        if company == platform:
            col = 3
            keys = {(_norm(getattr(e, "line", None)), r.section) for r in rows if not section or r.section == section}
        elif company == target:
            col, keys = 4, set()
            for t in target_lines:
                m = by_target.get(t.id)
                if m is not None and _line_matches(t, e):
                    key = _platform_row_of(platform_lines, m, t)
                    if key:
                        keys.add(key)
        else:
            continue
        for i, r in enumerate(rows):
            if (_norm(r.line), r.section) in keys:
                out[(i, col)] = e
    return out


def _bold(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)


def build_workbook(store, deal, result, items, chain_valid: bool, edits=None) -> Workbook:
    """The final workbook. `edits` defaults to the store's active edits for the deal (reverted ones are dropped)."""
    edits = cert_doc.deal_edits(store, deal.id, edits)
    cert = cert_doc.build_certificate(store, deal, result, items, chain_valid, edits=edits)

    wb = Workbook()
    ws = wb.active
    ws.title = "Combined"
    ws.append(["Line", "Section", deal.platform, deal.target, "Combined", "Run id"])
    for r in result.rows:
        ws.append([r.line, r.section, r.platform, r.target, r.combined, r.run_id])
        for cell in ws[ws.max_row][2:5]:
            cell.number_format = THOUSANDS
    ws.cell(row=1, column=4).comment = Comment(TARGET_EDIT_NOTE, "Stochastic")
    fill = PatternFill("solid", fgColor=AMBER)
    for (i, col), e in _edited_cells(store, deal, result.rows, edits).items():
        cell = ws.cell(row=i + 2, column=col)
        cell.fill = fill
        cell.comment = Comment(edit_note(e), "Stochastic")
    _bold(ws)

    we = wb.create_sheet("Edits")
    we.append(EDITS_HEADER)
    for e in edits:
        old, new = cert_doc.edit_values(e)
        we.append([getattr(e, "line", None), getattr(e, "company", None), old, new, cert_doc.edit_reason(e),
                   getattr(e, "actor", None), _when(getattr(e, "at", None)), getattr(e, "policy_id", None)])
        for cell in we[we.max_row][2:4]:  # an amount edit's agent and final values; a remap's are line names
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = THOUSANDS
    _bold(we)

    wa = wb.create_sheet("Audit")
    wa.append(AUDIT_HEADER)

    def tool_of(run_id):
        run = store.tool_runs.get(run_id) if run_id else None
        return run.tool if run else None

    def fig(name, block, key):
        if block:
            wa.append([name, block.get(key), block.get("run_id"), tool_of(block.get("run_id"))])
        else:
            wa.append([name, None, "no run", None])

    fig("EBITDA", cert["ebitda"], "ebitda")
    fig("Debt", cert["debt"], "debt")
    fig("Leverage", cert["leverage"], "leverage")
    lt = cert["limit_test"]
    fig("Limit test leverage", lt, "leverage")
    fig("Limit test limit", lt, "limit")
    fig("Limit test verdict", lt, "verdict")
    fig("Limit test headroom (turns)", lt, "headroom_turns")
    fig("Limit test debt capacity", lt, "debt_capacity")
    fig("Tie-out ok", cert["tie_out"], "ok")
    for r in result.rows:
        wa.append([f"Combined: {r.line} ({r.section})", r.combined, r.run_id, tool_of(r.run_id)])
    au = cert["approval"]["audit"]
    wa.append([None, None, None, None])
    wa.append(["Audit rows", au["rows"], None, None])
    wa.append(["Chain valid", au["chain_valid"], None, None])
    wa.append(["Denied attempts", au["denied"], None, None])
    wa.append(["Backend", cert["header"]["backend"],
               "Agent output backend. The gateway, the math, and the audit log are real.", None])
    wa.append(["Draft. Not signed.", None, None, None])
    _bold(wa)
    return wb


def workbook_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
